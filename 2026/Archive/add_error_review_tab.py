from openpyxl import load_workbook
import xlsxwriter, os

src='/mnt/data/Holy_Family_Parish_Survey_600dpi_Responses.xlsx'
out='/mnt/data/Holy_Family_Parish_Survey_600dpi_Responses_with_ERROR_Review.xlsx'

# Error review items identified in the existing workbook
review_items = [
    {
        'survey_id': 1,
        'pdf_pages': '1-2',
        'question': 'Q26 Mass attended',
        'current_value': 'ERROR: Sunday 9am and Sunday 7pm both appear marked',
        'issue': 'Single-choice question appears to have two Mass times marked.',
        'source_page': 'Page 1',
        'crop': '/mnt/data/q26_survey1_crop.png'
    },
    {
        'survey_id': 3,
        'pdf_pages': '5-6',
        'question': 'Q26 Mass attended',
        'current_value': 'ERROR: Sunday 9am and Sunday 7pm both appear marked',
        'issue': 'Single-choice question appears to have two Mass times marked.',
        'source_page': 'Page 5',
        'crop': '/mnt/data/q26_survey3_crop.png'
    },
]

wb_src=load_workbook(src, data_only=True)
# collect sheet data except if existing error sheet
sheets=[]
for ws in wb_src.worksheets:
    if ws.title == 'ERROR review':
        continue
    data=[[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    sheets.append((ws.title, data))

workbook=xlsxwriter.Workbook(out)
header_fmt=workbook.add_format({'bold':True,'bg_color':'#1F4E78','font_color':'white','text_wrap':True,'valign':'top'})
error_fmt=workbook.add_format({'bg_color':'#F4CCCC','text_wrap':True,'valign':'top'})
review_fmt=workbook.add_format({'bg_color':'#FCE4D6','text_wrap':True,'valign':'top'})
wrap_fmt=workbook.add_format({'text_wrap':True,'valign':'top'})

for title,data in sheets:
    ws=workbook.add_worksheet(title[:31])
    if not data: continue
    nrows=len(data); ncols=max(len(r) for r in data)
    for r,row in enumerate(data):
        for c,val in enumerate(row):
            fmt=wrap_fmt
            if r==0: fmt=header_fmt
            elif isinstance(val,str) and 'ERROR' in val: fmt=error_fmt
            elif isinstance(val,str) and ('UNCLEAR' in val or 'Review recommended' in val): fmt=review_fmt
            ws.write(r,c,val,fmt)
    if title=='Digitised responses':
        ws.freeze_panes(1,4); ws.autofilter(0,0,nrows-1,ncols-1)
    for c in range(ncols):
        width=20
        if title=='Digitised responses':
            width=18
            if c in [3,39]: width=45
            if c in [6,15,22,31]: width=35
            if c in [0,1,2]: width=16
        elif title=='OCR text by page' and c==2:
            width=120
        elif c==1:
            width=110 if title=='Method and review notes' else 25
        elif c==0:
            width=45
        ws.set_column(c,c,width)

# Add review tab with crops
ws=workbook.add_worksheet('ERROR review')
cols=['Survey ID','PDF pages','Question','Current value','Reason for review','Source page','Cropped image']
for c,h in enumerate(cols): ws.write(0,c,h,header_fmt)
ws.set_column(0,0,10); ws.set_column(1,1,14); ws.set_column(2,2,24); ws.set_column(3,4,42); ws.set_column(5,5,12); ws.set_column(6,6,80)
for r,item in enumerate(review_items, start=1):
    ws.set_row(r, 230)
    ws.write(r,0,item['survey_id'],wrap_fmt)
    ws.write(r,1,item['pdf_pages'],wrap_fmt)
    ws.write(r,2,item['question'],wrap_fmt)
    ws.write(r,3,item['current_value'],error_fmt)
    ws.write(r,4,item['issue'],wrap_fmt)
    ws.write(r,5,item['source_page'],wrap_fmt)
    if os.path.exists(item['crop']):
        ws.insert_image(r,6,item['crop'],{'x_scale':0.38,'y_scale':0.38,'x_offset':4,'y_offset':4})
    else:
        ws.write(r,6,'Crop image not found',error_fmt)
ws.freeze_panes(1,0)

workbook.close()
print(out)
